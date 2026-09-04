"""台账导入：模板下载（列定义与解析共用常量）+ xlsx 上传覆盖 + 导入记录"""

import io
import time

from fastapi import APIRouter, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel

from app.db.mongo_store import store
from app.errors import BusinessError
from app.services.pagination import paged_query

router = APIRouter(prefix='/api/import', tags=['import'])

TEMPLATE_COLUMNS = {
    'commercial': ['签约日期', '经营单元', '行业', '产品线', '产品型号', '合同金额(万元)', '收入额(万元)', '客户'],
    'ppl': ['年份', '项目名称', '经营单元', '行业', '合同金额(万元)', '阶段', '风险等级'],
    'goal': ['年份', '经营单元', '商业目标(万元)', '商解目标(万元)'],
}
_SAMPLE_ROW = {
    'commercial': ['2026-06-15', '北京代表处', '政企', '智能计算', 'JS-9100', 800, 640, '华信集团'],
    'ppl': [2026, '北京智能计算项目1', '北京代表处', '政企', 1200, '商机', '高'],
    'goal': [2026, '北京代表处', 5000, 1800],
}
_TYPES = set(TEMPLATE_COLUMNS)


class _SkipRow(Exception):
    pass


@router.get('/template')
async def template(type: str = 'commercial'):
    if type not in _TYPES:
        raise BusinessError('type 参数不合法', 400)
    wb = Workbook()
    ws = wb.active
    ws.title = '模板'
    cols = TEMPLATE_COLUMNS[type]
    ws.append(cols)
    ws.append(_SAMPLE_ROW[type])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={type}_template.xlsx'},
    )


@router.post('/upload')
async def upload(type: str, year: int, file: UploadFile):
    if type not in _TYPES:
        raise BusinessError('type 参数不合法', 400)
    cols = TEMPLATE_COLUMNS[type]
    try:
        wb = load_workbook(io.BytesIO(await file.read()), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        await _log(type, year, file.filename, 0, 0, '失败-格式不符', '文件解析失败')
        raise BusinessError('xlsx 解析失败', 400)
    if not rows or list(rows[0])[:len(cols)] != cols:
        await _log(type, year, file.filename, max(0, len(rows) - 1), 0, '失败-格式不符', '表头与模板不一致')
        raise BusinessError('表头与模板不一致，请下载模板填写', 400)

    docs: list[dict] = []
    errors: list[str] = []
    for i, row in enumerate(rows[1:]):
        if row is None or all(v is None for v in row):
            continue
        try:
            docs.append(_parse_row(type, year, row))
        except Exception as e:
            errors.append(f'第{i + 2}行: {e}')
    status = '成功' if not errors and docs else ('部分成功' if docs else '失败-格式不符')
    if docs:
        await store.remove(type_collection(type), {'year': year})
        await store.insert_many(type_collection(type), docs)
    await _log(type, year, file.filename or '', len(rows) - 1, len(docs), status, '; '.join(errors[:5]))
    return {'ok': status != '失败-格式不符', 'status': status, 'total': max(0, len(rows) - 1),
            'success': len(docs), 'errors': errors[:10]}


class ImportLogIn(BaseModel):
    pass


def type_collection(t: str) -> str:
    return {'commercial': 'CommercialLedger', 'ppl': 'PplLedger', 'goal': 'GoalLedger'}[t]


def _parse_row(t: str, year: int, row) -> dict:
    def s(v) -> str:
        return str(v).strip() if v is not None else ''

    def f(v) -> float:
        return float(v or 0)

    if t == 'commercial':
        sign_date = s(row[0]) or f'{year}-01-01'
        month = int(sign_date[5:7]) if len(sign_date) >= 7 and sign_date[5:7].isdigit() else 1
        return {'signDate': sign_date, 'year': int(sign_date[:4]) if sign_date[:4].isdigit() else year,
                'quarter': (month - 1) // 3 + 1, 'month': month,
                'unit': s(row[1]), 'industry': s(row[2]), 'productLine': s(row[3]),
                'productModel': s(row[4]), 'contractAmt': f(row[5]), 'income': f(row[6]),
                'customer': s(row[7])}
    if t == 'ppl':
        return {'year': int(row[0] or year), 'projectName': s(row[1]), 'unit': s(row[2]),
                'industry': s(row[3]), 'contractAmt': f(row[4]), 'stage': s(row[5]),
                'riskLevel': s(row[6]) or '中'}
    return {'year': int(row[0] or year), 'unit': s(row[1]),
            'commercialGoal': f(row[2]), 'solutionGoal': f(row[3])}


async def _log(t: str, year: int, fn: str, total: int, ok: int, status: str, detail: str) -> None:
    await store.insert('ImportLog', {'type': t, 'year': year, 'fileName': fn, 'totalRows': total,
                                     'successRows': ok, 'status': status, 'detail': detail[:500]})


@router.get('/log')
async def import_log(page: int = 1, pageSize: int = 10, type: str = ''):
    cond = {'type': type} if type else {}
    return await paged_query('ImportLog', cond, page, pageSize,
                             '_id, type, year, fileName, totalRows, successRows, status, detail, userName, createdAt')
